"""Read-only, bounded previews for XLSX, Arrow IPC/Feather, JSON Lines and SQLite.

Usage: tables_helper.py summarize <path> [sheet-or-table] [unused-slice]
Exit codes: 0 ok, 3 missing dependency, 4 missing file, 5 invalid/oversized input.
"""
from __future__ import annotations

import itertools
import json
import math
import sqlite3
import sys
import zipfile
from pathlib import Path

MAX_ROWS = 200
MAX_COLS = 50
MAX_CELL = 512
MAX_COLLECTIONS = 100
MAX_BYTES = 64 * 1024 * 1024


def need(module):
    try:
        __import__(module)
    except ImportError:
        sys.stderr.write(f"{module} not installed; restart the app to finish preview setup\n")
        sys.exit(3)


def cell(value):
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, int):
        # JSON numbers would silently round larger identifiers in the browser.
        return str(value) if abs(value) > 2**53 - 1 else value
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, bytes):
        return f"[binary: {len(value)} bytes]"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, default=str)[:MAX_CELL]
    return str(value)[:MAX_CELL]


def table(columns, rows, total=None, num_columns=None, **extra):
    more = len(rows) > MAX_ROWS
    head = [[cell(v) for v in row[:MAX_COLS]] for row in rows[:MAX_ROWS]]
    return {
        "kind": "table", "columns": columns[:MAX_COLS], "head": head,
        "num_rows": total, "num_columns": num_columns if num_columns is not None else len(columns),
        "rows_truncated": more or (total is not None and total > len(head)),
        "columns_truncated": (num_columns or len(columns)) > MAX_COLS,
        "cell_limit": MAX_CELL, **extra,
    }


def choose(names, selected):
    if selected and selected not in names:
        raise ValueError("The selected sheet or table is no longer available")
    return selected or next(iter(names), "")


def xlsx(path, selected):
    need("openpyxl")
    from openpyxl import load_workbook

    with zipfile.ZipFile(path) as archive:
        if len(archive.infolist()) > 10_000 or sum(i.file_size for i in archive.infolist()) > MAX_BYTES:
            raise ValueError("Workbook exceeds the 64 MiB expanded preview limit")
    book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        names = book.sheetnames[:MAX_COLLECTIONS]
        key = choose(names, selected)
        if not key:
            return table([], [], 0, collections=[], selected="")
        sheet = book[key]
        declared_cols = sheet.max_column or 0
        # Producer-supplied worksheet dimensions can be incorrect. Bound actual
        # iteration ourselves rather than trusting a bogus A1:A1 declaration.
        sheet.reset_dimensions()
        values = list(itertools.islice(sheet.iter_rows(max_col=MAX_COLS, values_only=True), MAX_ROWS + 2))
        used = max((i + 1 for row in values for i, v in enumerate(row) if v is not None), default=0)
        width = min(MAX_COLS, max(declared_cols, used))
        headers = values[0][:width] if values else []
        columns = [{"name": str(v)[:MAX_CELL] if v is not None else f"Column {i + 1}", "dtype": "cell"}
                   for i, v in enumerate(headers)]
        rows = [list(row[:width]) for row in values[1:]]
        return table(columns, rows, len(rows) if len(rows) <= MAX_ROWS else None,
                     num_columns=max(declared_cols, width), collections=names, selected=key,
                     collections_truncated=len(book.sheetnames) > MAX_COLLECTIONS,
                     note="First row supplies column names. Formulas show saved values only; formulas are not recalculated.")
    finally:
        book.close()


def jsonlines(path, selected):
    if selected:
        raise ValueError("JSON Lines has no sheets or tables to select")
    records = []
    names = []
    seen = set()
    bytes_read = 0
    complete = False
    with path.open("rb") as stream:
        for line_number in range(100_000):
            raw = stream.readline(1024 * 1024 + 1)
            if not raw:
                complete = True
                break
            bytes_read += len(raw)
            if len(raw) > 1024 * 1024 or bytes_read > 8 * 1024 * 1024:
                raise ValueError("JSON Lines preview exceeds its 1 MiB line / 8 MiB scan limit")
            if not raw.strip():
                continue
            try:
                value = json.loads(raw.decode("utf-8-sig"))
            except (ValueError, UnicodeError) as exc:
                raise ValueError(f"Invalid JSON on line {line_number + 1}: {exc}") from exc
            if not isinstance(value, dict):
                value = {"value": value}
            # Only infer the schema from the displayed sample.
            if len(records) < MAX_ROWS:
                for name in value:
                    if name not in seen:
                        seen.add(name)
                        names.append(name)
            records.append(value)
            if len(records) > MAX_ROWS:
                break
    columns = [{"name": name[:MAX_CELL], "dtype": "JSON"} for name in names[:MAX_COLS]]
    rows = [[record.get(name) for name in names[:MAX_COLS]] for record in records]
    return table(columns, rows, len(records) if complete else None, num_columns=len(names),
                 note="Columns are inferred from the displayed records; nested values are shown as JSON.")


def sqlite(path, selected):
    # URI escaping comes from as_uri (including # and ? in filenames).
    connection = sqlite3.connect(path.resolve().as_uri() + "?mode=ro", uri=True, timeout=2)
    try:
        connection.execute("PRAGMA query_only=ON")
        connection.execute("PRAGMA trusted_schema=OFF")
        connection.setlimit(sqlite3.SQLITE_LIMIT_LENGTH, 2 * 1024 * 1024)
        steps = 0

        def progress():
            nonlocal steps
            steps += 1
            return int(steps > 2000)

        connection.set_progress_handler(progress, 1000)
        # Do not execute views or virtual tables defined by an uploaded file.
        entries = connection.execute(
            "SELECT name FROM sqlite_schema WHERE type='table' AND name NOT GLOB 'sqlite_*' "
            "AND upper(ltrim(sql)) NOT LIKE 'CREATE VIRTUAL%' ORDER BY name LIMIT ?",
            (MAX_COLLECTIONS + 1,),
        ).fetchall()
        names = [r[0] for r in entries[:MAX_COLLECTIONS]]
        key = choose(names, selected)
        if not key:
            return table([], [], 0, collections=[], selected="", note="No ordinary tables to preview.")
        quote = lambda s: '"' + s.replace('"', '""') + '"'
        # table_info omits generated columns. table_xinfo marks generated
        # columns as 2/3, while 1 denotes hidden virtual-table fields.
        fields = [field for field in connection.execute(f"PRAGMA table_xinfo({quote(key)})").fetchall()
                  if field[6] != 1]
        columns = [{"name": field[1], "dtype": field[2] or "dynamic"} for field in fields[:MAX_COLS]]
        expressions = []
        for col in columns:
            ident = quote(col["name"])
            expressions.append(
                f"CASE typeof({ident}) WHEN 'blob' THEN '[binary: ' || length({ident}) || ' bytes]' "
                f"WHEN 'text' THEN substr({ident}, 1, {MAX_CELL}) ELSE {ident} END"
            )
        rows = connection.execute(f"SELECT {', '.join(expressions)} FROM {quote(key)} LIMIT ?", (MAX_ROWS + 1,)).fetchall()
        return table(columns, rows, len(rows) if len(rows) <= MAX_ROWS else None, num_columns=len(fields),
                     collections=names, selected=key, collections_truncated=len(entries) > MAX_COLLECTIONS,
                     note="Read-only sample of ordinary tables. Views and virtual tables are excluded.")
    finally:
        connection.close()


def arrow(path, selected):
    if selected:
        raise ValueError("Arrow has no sheets or tables to select")
    need("pyarrow")
    import pyarrow as pa
    if path.stat().st_size > MAX_BYTES:
        raise ValueError("Arrow preview currently supports files up to 64 MiB")
    with pa.memory_map(str(path), "r") as source:
        try:
            reader = pa.ipc.open_file(source)
            batches = (reader.get_batch(i) for i in range(reader.num_record_batches))
        except pa.ArrowInvalid:
            source.seek(0)
            reader = pa.ipc.open_stream(source)
            batches = iter(reader)
        schema = reader.schema
        columns = [{"name": field.name, "dtype": str(field.type)} for field in list(schema)[:MAX_COLS]]
        rows = []
        complete = True
        for batch in batches:
            if batch.nbytes > MAX_BYTES:
                raise ValueError("Arrow record batch exceeds the 64 MiB preview limit")
            batch = batch.slice(0, MAX_ROWS + 1 - len(rows)).select(list(range(min(len(schema), MAX_COLS))))
            values = [col.to_pylist() for col in batch.columns]
            rows.extend(list(row) for row in zip(*values))
            if len(rows) > MAX_ROWS:
                complete = False
                break
        return table(columns, rows, len(rows) if complete else None, num_columns=len(schema),
                     note="Arrow IPC file/stream and Feather v2. Charts use the displayed row sample.")


DISPATCH = {"xlsx": xlsx, "arrow": arrow, "feather": arrow, "ipc": arrow,
            "jsonl": jsonlines, "ndjson": jsonlines, "sqlite": sqlite, "sqlite3": sqlite, "db": sqlite}


def main():
    try:
        if len(sys.argv) < 3 or sys.argv[1] != "summarize":
            raise ValueError("usage: tables_helper.py summarize <path> [sheet-or-table]")
        path = Path(sys.argv[2])
        if not path.is_file():
            sys.stderr.write("File not found\n")
            sys.exit(4)
        ext = path.suffix.lower().lstrip(".")
        if ext not in DISPATCH:
            raise ValueError(f"Unsupported extension: {ext}")
        result = DISPATCH[ext](path, sys.argv[3] if len(sys.argv) > 3 else "")
        result.update(format=ext, file_size=path.stat().st_size)
        sys.stdout.write(json.dumps(result, allow_nan=False))
    except Exception as exc:
        sys.stderr.write(f"{type(exc).__name__}: {exc}\n")
        sys.exit(5)


if __name__ == "__main__":
    main()
