"""Regenerate committed Office integration fixtures:
uv run --with python-docx --with python-pptx --with openpyxl --with pillow generate.py
No generator dependencies are needed to run the normal test suite.
"""
from pathlib import Path
from io import BytesIO
from docx import Document
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as SlideInches
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from PIL import Image

out = Path(__file__).parent
image = BytesIO()
Image.new('RGB', (240, 100), '#6478bb').save(image, format='PNG')
image.seek(0)
doc = Document()
doc.add_heading('Research overview', 0)
p = doc.add_paragraph()
p.add_run('Baseline ').bold = True
p.add_run('findings').italic = True
doc.add_paragraph('A paragraph with unchanged formatting.')
t = doc.add_table(rows=2, cols=2)
t.cell(0, 0).text = 'Condition'
t.cell(0, 1).text = 'Value'
t.cell(1, 0).text = 'Control'
t.cell(1, 1).text = '12'
doc.add_picture(image, width=Inches(2.5))
doc.sections[0].header.paragraphs[0].text = 'Study header'
doc.save(out / 'report.docx')
pres = Presentation()
slide = pres.slides.add_slide(pres.slide_layouts[5])
slide.shapes.title.text = 'Research overview'
box = slide.shapes.add_textbox(SlideInches(1), SlideInches(2), SlideInches(7), SlideInches(1))
p = box.text_frame.paragraphs[0]
r = p.add_run(); r.text = 'Baseline '; r.font.bold = True
r = p.add_run(); r.text = 'findings'; r.font.italic = True
image.seek(0)
slide.shapes.add_picture(image, SlideInches(1), SlideInches(4), width=SlideInches(3))
slide.notes_slide.notes_text_frame.text = 'Speaker notes'
slide = pres.slides.add_slide(pres.slide_layouts[5])
slide.shapes.title.text = 'Results'
t = slide.shapes.add_table(2, 2, SlideInches(1), SlideInches(2), SlideInches(6), SlideInches(2)).table
t.cell(0, 0).text = 'Condition'; t.cell(0, 1).text = 'Value'
t.cell(1, 0).text = 'Control'; t.cell(1, 1).text = '12'
pres.save(out / 'slides.pptx')
book = Workbook()
sheet = book.active; sheet.title = 'Measurements'
sheet.append(['Condition', 'Value', 'Formula', 'Literal', 'Boolean'])
sheet.append(['Control', 12, '=B2*2', '=not a formula', True]); sheet['D2'].data_type = 's'
sheet['B2'].number_format = '0.00'
sheet['A4'] = 'Merged title'; sheet.merge_cells('A4:C4')
sheet['A51'] = 'Later page'
chart = BarChart(); chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2), titles_from_data=True)
sheet.add_chart(chart, 'G3')
second = book.create_sheet('Second sheet'); second.append(['Other', 99])
locked = book.create_sheet('Protected'); locked['A1'] = 'Locked'; locked.protection.sheet = True
book.save(out / 'workbook.xlsx')
